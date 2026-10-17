"""导出路由的分支选择：编辑过的 HTML / 模板回填 / Markdown。

分支选错的症状是"下载下来不是我看到的"，而且要等打开文件才发现。所以这里
真的调接口、真的把文件读回来验。
"""

from __future__ import annotations

import uuid

import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.services import draft_service, user_service

API = "/api/v1"
PASSWORD = "Test-Pa55word!"
TPL = "广东_劳动仲裁申请书"

EDITED = ('<p style="text-align:center">'
          '<span style="font-family:\'黑体\';font-size:16pt;font-weight:700">'
          '劳动仲裁申请书</span></p>'
          '<p style="text-indent:24.1pt"><span style="font-size:12pt">'
          '申请人于 2024 年 1 月入职。</span></p>')


def _uniq(p: str) -> str:
    return f"{p}_{uuid.uuid4().hex[:8]}"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def token(client: TestClient) -> str:
    team = user_service.create_team(_uniq("导出测试组"))
    u = user_service.create_user(_uniq("exporter"), PASSWORD, role="lawyer",
                                 full_name="导出测试律师", team_id=team["id"])
    r = client.post(f"{API}/auth/login",
                    json={"username": u["username"], "password": PASSWORD})
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture
def made():
    ids: list[str] = []
    yield ids
    for did in ids:
        draft_service.delete_draft(did)


def _export(client, token, did, fmt, body=None):
    r = client.post(f"{API}/drafts/{did}/export", headers={"Authorization": f"Bearer {token}"},
                    params={"fmt": fmt}, json=body or {})
    assert r.status_code == 200, f"导出失败：{r.status_code} {r.text}"
    return r.json()


def _download(client, token, did, tmp_path, suffix):
    r = client.get(f"{API}/drafts/{did}/download",
                   headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 200, r.text
    p = tmp_path / f"got{suffix}"
    p.write_bytes(r.content)
    return p


def test_edited_html_wins_over_template(client: TestClient, token: str, made,
                                        tmp_path) -> None:
    """存了 ``content_html`` 就导出这一份，**不再回填模板**。

    这是"所见即所得"的核心：律师在编辑页加了粗，下载的 Word 就该是粗的。
    如果这里选错分支，导出的会是"原始模板 + fills"，手动排版全丢。
    """
    did = client.post(f"{API}/drafts", headers={"Authorization": f"Bearer {token}"},
                      json={"content": "纯文本正文", "doc_type": "arbitration_application",
                            "title": _uniq("编辑过"), "template_id": TPL,
                            "fills": {"p1": "模板里填的值"}}).json()["draft_id"]
    made.append(did)
    client.put(f"{API}/drafts/{did}", headers={"Authorization": f"Bearer {token}"},
               json={"content": "纯文本正文", "content_html": EDITED})

    info = _export(client, token, did, "docx")
    doc = Document(str(_download(client, token, did, tmp_path, ".docx")))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    joined = "\n".join(p.text for p in paras)

    assert "申请人于 2024 年 1 月入职。" in joined, f"没导出编辑后的正文：{joined!r}"
    assert "模板里填的值" not in joined, "回填了模板，手动编辑的内容被顶掉了"

    title = paras[0]
    assert title.text == "劳动仲裁申请书"
    assert title.runs[0].font.bold is True, "加粗没导出"
    assert title.runs[0].font.size.pt == 16, f"字号没导出：{title.runs[0].font.size}"
    assert "CENTER" in str(title.alignment), "居中没导出"
    body = [p for p in paras if p.text.startswith("申请人于")][0]
    assert body.paragraph_format.first_line_indent is not None, "首行缩进没导出"


def test_client_supplied_html_beats_stored(client: TestClient, token: str, made,
                                           tmp_path) -> None:
    """请求里带了 ``content_html`` 就用请求里的 —— 处理"还没保存就下载"。"""
    did = client.post(f"{API}/drafts", headers={"Authorization": f"Bearer {token}"},
                      json={"content": "正文", "doc_type": "arbitration_application",
                            "title": _uniq("未保存")}).json()["draft_id"]
    made.append(did)
    client.put(f"{API}/drafts/{did}", headers={"Authorization": f"Bearer {token}"},
               json={"content": "正文", "content_html": "<p>存库的版本</p>"})

    info = _export(client, token, did, "docx",
                   {"content_html": "<p>屏幕上还没保存的版本</p>"})
    doc = Document(str(_download(client, token, did, tmp_path, ".docx")))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "屏幕上还没保存的版本" in joined, joined
    assert "存库的版本" not in joined


def test_exported_html_is_sanitized(client: TestClient, token: str, made,
                                    tmp_path) -> None:
    """导出时也要过净化器 —— 这是同一条不可信输入的另一条出口。"""
    did = client.post(f"{API}/drafts", headers={"Authorization": f"Bearer {token}"},
                      json={"content": "正文", "doc_type": "arbitration_application",
                            "title": _uniq("带脚本")}).json()["draft_id"]
    made.append(did)

    _export(client, token, did, "docx",
            {"content_html": "<p>正常文字</p><script>alert(1)</script>"})
    doc = Document(str(_download(client, token, did, tmp_path, ".docx")))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "正常文字" in joined, joined
    assert "script" not in joined.lower(), f"脚本被写进 Word 了：{joined!r}"


def test_falls_back_to_template_when_not_edited(client: TestClient, token: str, made,
                                                tmp_path) -> None:
    """没编辑过就走模板回填 —— 这条是原有能力，别被新分支挤掉。"""
    from backend.app.services import template_service as ts

    if not (ts.template_dir() / f"{TPL}.docx").exists():
        pytest.skip(f"模板 {TPL} 不在模板目录")

    did = client.post(f"{API}/drafts", headers={"Authorization": f"Bearer {token}"},
                      json={"content": "正文", "doc_type": "arbitration_application",
                            "title": _uniq("未编辑"), "template_id": TPL,
                            "fills": {"p1": "申请人：李四，男。"}}).json()["draft_id"]
    made.append(did)

    _export(client, token, did, "docx")
    doc = Document(str(_download(client, token, did, tmp_path, ".docx")))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "申请人：李四，男。" in joined, f"模板回填没生效：{joined[:200]!r}"


def test_pdf_export_of_edited_html(client: TestClient, token: str, made, tmp_path) -> None:
    """编辑过的内容导出 PDF 也要成功（不能因为走了新分支就报错）。"""
    did = client.post(f"{API}/drafts", headers={"Authorization": f"Bearer {token}"},
                      json={"content": "正文", "doc_type": "arbitration_application",
                            "title": _uniq("PDF")}).json()["draft_id"]
    made.append(did)
    client.put(f"{API}/drafts/{did}", headers={"Authorization": f"Bearer {token}"},
               json={"content": "正文", "content_html": EDITED})

    try:
        _export(client, token, did, "pdf")
    except AssertionError as exc:
        if "中文字体" in str(exc):
            pytest.skip("本机没装中文字体，PDF 导出不可用")
        raise
    p = _download(client, token, did, tmp_path, ".pdf")
    assert p.read_bytes()[:4] == b"%PDF", "导出的不是 PDF"

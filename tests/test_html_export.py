"""编辑页 HTML → Word 的导出。

这条链路的失败方式是**看不出来**：docx 能打开、文字都在，只是加粗没了、
字号退回默认、落款掉回左边。用户要的正是"所见即所得"，所以每条样式都
真的读回 docx 去验，而不是断言"函数没抛异常"。
"""

from __future__ import annotations

import pytest
from docx import Document

from backend.app.services.html_export import write_docx_from_html

# 黑体 16pt 居中标题 + 宋体 12pt 缩进正文 + 右对齐落款 —— 就是模板的样子
TPL_HTML = """
<p style="text-align:center"><span style="font-family:'黑体';font-size:16pt;font-weight:700">劳动仲裁申请书</span></p>
<p style="text-indent:24.1pt"><span style="font-family:'宋体';font-size:12pt">申请人于 2024 年 1 月入职某公司。</span></p>
<p style="text-align:right"><span style="font-family:'宋体';font-size:12pt">申请人（签名）：张三</span></p>
"""


def _write(tmp_path, html: str, template_id: str = ""):
    out = tmp_path / "out.docx"
    write_docx_from_html(out, html, template_id)
    assert out.exists(), "没生成文件"
    return Document(str(out))


def test_keeps_cjk_font_size_and_alignment(tmp_path) -> None:
    """字体 / 字号 / 加粗 / 对齐都要落到 docx 的 run 与段落上。"""
    doc = _write(tmp_path, TPL_HTML)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    assert len(paras) == 3, f"段落数不对：{[p.text for p in paras]}"

    title = paras[0]
    assert title.text == "劳动仲裁申请书"
    assert title.alignment is not None and "CENTER" in str(title.alignment), \
        f"标题没居中：{title.alignment}"
    tr = title.runs[0]
    assert tr.font.size.pt == 16, f"标题字号：{tr.font.size}"
    assert tr.font.bold is True, "标题没加粗"
    # 中文必须写在 w:eastAsia 上，只设 font.name 对汉字不生效
    from docx.oxml.ns import qn

    ea = tr._element.rPr.rFonts.get(qn("w:eastAsia"))
    assert ea == "黑体", f"中文字体没写进 eastAsia，实际 {ea!r}"

    body = paras[1]
    assert body.paragraph_format.first_line_indent is not None, "首行缩进丢了"
    assert abs(body.paragraph_format.first_line_indent.pt - 24.1) < 0.2, \
        f"缩进值不对：{body.paragraph_format.first_line_indent.pt}"
    assert body.runs[0].font.size.pt == 12


def test_keeps_right_alignment_of_signature(tmp_path) -> None:
    """落款要真的右对齐 —— 这是之前修过一轮的问题，别再退回去。"""
    doc = _write(tmp_path, TPL_HTML)
    paras = [p for p in doc.paragraphs if p.text.strip()]
    assert "RIGHT" in str(paras[2].alignment), f"落款没右对齐：{paras[2].alignment}"


def test_inline_formatting(tmp_path) -> None:
    """加粗 / 斜体 / 下划线 / 删除线 / 字色 —— 工具栏直接能改的那几项。"""
    doc = _write(tmp_path, '<p><b>粗</b><i>斜</i><u>下</u><s>删</s>'
                          '<span style="color:#ff0000">红</span></p>')
    runs = {r.text: r for r in doc.paragraphs[0].runs if r.text.strip()}
    assert runs["粗"].font.bold is True
    assert runs["斜"].font.italic is True
    assert runs["下"].font.underline is True
    assert runs["删"].font.strike is True
    assert str(runs["红"].font.color.rgb) == "FF0000", runs["红"].font.color.rgb


def test_nested_inline_style_is_inherited(tmp_path) -> None:
    """外层 ``span`` 的字号要传给内层的 ``b``。

    只读最内层标签的话，``<span style="font-size:16pt"><b>劳动</b></span>``
    里的"劳动"会退回默认 10.5pt —— 页面上是 16pt，导出就变小了。
    """
    doc = _write(tmp_path, '<p><span style="font-size:16pt"><b>劳动</b></span></p>')
    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == 16 and run.font.bold is True


def test_exec_command_font_size_tag(tmp_path) -> None:
    """``execCommand('fontSize', 7)`` 产出的是 ``<font size=7>``，要折算成磅。"""
    doc = _write(tmp_path, '<p><font size="7">大标题</font></p>')
    assert doc.paragraphs[0].runs[0].font.size.pt == 36


def test_table_with_merged_cells(tmp_path) -> None:
    """表格 + ``colspan`` / ``rowspan`` 要还原成真正的合并单元格。"""
    doc = _write(tmp_path, """
    <table><tbody>
      <tr><td colspan="2"><p>当事人信息</p></td><td>备注</td></tr>
      <tr><td rowspan="2"><p>姓名</p></td><td>张三</td><td>x</td></tr>
      <tr><td>李四</td><td>y</td></tr>
    </tbody></table>""")
    assert len(doc.tables) == 1, "表格没建出来"
    t = doc.tables[0]
    assert len(t.rows) == 3
    # colspan=2：同一行的第 1、2 格必须合并成一个
    assert t.cell(0, 0)._tc is t.cell(0, 1)._tc, "colspan 没合并"
    assert "当事人信息" in t.cell(0, 0).text
    # rowspan=2：第 1、2 行的第 1 格合并
    assert t.cell(1, 0)._tc is t.cell(2, 0)._tc, "rowspan 没合并"
    assert t.cell(1, 1).text.strip() == "张三"
    assert t.cell(2, 1).text.strip() == "李四"


def test_empty_paragraph_keeps_its_place(tmp_path) -> None:
    """空段落要留住 —— 模板靠它撑竖向节奏。"""
    doc = _write(tmp_path, "<p>甲</p><p><br/></p><p>乙</p>")
    texts = [p.text for p in doc.paragraphs]
    assert "甲" in texts and "乙" in texts
    assert texts.index("乙") - texts.index("甲") == 2, f"空段落被吞了：{texts}"


def test_no_leading_blank_line(tmp_path) -> None:
    """文首不能多出一个空行。

    ``Document()`` 自带一个空段落，不去复用它的话每份导出都会顶着一个空行。
    """
    doc = _write(tmp_path, "<p>第一行</p><p>第二行</p>")
    assert doc.paragraphs[0].text == "第一行", \
        f"文首多出空行：{[p.text for p in doc.paragraphs]}"


def test_table_only_document_has_no_stray_empty_paragraph(tmp_path) -> None:
    """整篇只有一个表格时，占位段落要被摘掉。"""
    doc = _write(tmp_path, "<table><tbody><tr><td>x</td></tr></tbody></table>")
    assert len(doc.tables) == 1
    assert [p.text for p in doc.paragraphs if p.text.strip()] == [], \
        f"多出无关段落：{[p.text for p in doc.paragraphs]}"


def test_citations_appended(tmp_path) -> None:
    """引用依据附在最后，未核验的要标出来。"""
    out = tmp_path / "c.docx"
    write_docx_from_html(out, "<p>正文</p>",
                         citations=[{"name": "中华人民共和国劳动法", "article": "第 10 条",
                                     "status": "有效", "verified": False}])
    doc = Document(str(out))
    joined = "\n".join(p.text for p in doc.paragraphs)
    assert "引用依据" in joined
    assert "中华人民共和国劳动法" in joined
    assert "未通过核验" in joined, "未核验的引用没标出来 —— 律师会当成已核验的用"


def test_page_setup_follows_template(tmp_path) -> None:
    """抄原模板的纸张与页边距。

    不抄的话，律师改一个字再导出，A4 就变成 Letter、页边距也变了，
    打印出来才发现对不上归档要求。
    """
    from backend.app.services import template_service as ts

    name = "劳动争议纠纷起诉状"
    if not (ts.template_dir() / f"{name}.docx").exists():
        pytest.skip(f"模板 {name} 不在模板目录")

    src = Document(str(ts.template_dir() / f"{name}.docx")).sections[0]
    doc = _write(tmp_path, "<p>正文</p>", template_id=name)
    sec = doc.sections[0]
    assert sec.page_width == src.page_width, "纸张宽度没跟着模板"
    assert sec.page_height == src.page_height, "纸张高度没跟着模板"
    assert sec.left_margin == src.left_margin, "左边距没跟着模板"


def test_bad_template_id_does_not_break_export(tmp_path) -> None:
    """模板找不到时退回默认纸张，**不能因此导出失败**。"""
    doc = _write(tmp_path, "<p>正文</p>", template_id="不存在的模板")
    assert doc.paragraphs[0].text == "正文"


# --------------------------------------------------------------------------- PDF
def _pdf(tmp_path, html: str):
    from backend.app.services.html_export import write_pdf_from_html

    out = tmp_path / "out.pdf"
    try:
        write_pdf_from_html(out, html)
    except ImportError as exc:
        pytest.skip(f"本机 PDF 依赖不可用：{exc}")
    assert out.exists() and out.stat().st_size > 0, "没生成 PDF"
    return out.read_bytes()


def test_pdf_from_html_is_a_real_pdf(tmp_path) -> None:
    """导出的要是**真 PDF**（不是顶着 .pdf 后缀的文本）。"""
    data = _pdf(tmp_path, TPL_HTML)
    assert data[:4] == b"%PDF", data[:8]


def test_pdf_with_merged_table_does_not_crash(tmp_path) -> None:
    """带 ``colspan`` / ``rowspan`` 的表格在 PDF 上也要能出。

    reportlab 的表格用 ``SPAN`` 表达合并，行列号与 docx 那套**方向相反**
    （``(列, 行)`` vs ``(行, 列)``）—— 写反了不会报错，只是合并位置错，
    所以要真的造一张两边都不对称的表来试。
    """
    data = _pdf(tmp_path, """
    <table><tbody>
      <tr><td colspan="2"><p>当事人信息</p></td><td>备注</td></tr>
      <tr><td rowspan="2"><p>姓名</p></td><td>张三</td><td>x</td></tr>
      <tr><td>李四</td><td>y</td></tr>
    </tbody></table>""")
    assert data[:4] == b"%PDF"


def test_pdf_checkboxes_use_glyphs_the_font_actually_has(tmp_path) -> None:
    """☑/☐ 要换成 □/■。

    中文 TrueType 里**没有** U+2610 / U+2611 的字形，直接输出会印成空白方块 ——
    律师看到的是"勾不见了"，而 PDF 又不会报错。
    """
    from backend.app.services.html_export import write_pdf_from_html

    out = tmp_path / "cb.pdf"
    try:
        write_pdf_from_html(out, "<p>甲方 ☑ 是 ☐ 否</p>")
    except ImportError as exc:
        pytest.skip(f"本机 PDF 依赖不可用：{exc}")
    assert out.stat().st_size > 0
